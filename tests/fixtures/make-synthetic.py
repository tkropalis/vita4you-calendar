"""Builds tests/fixtures/synthetic.xlsx: a fictional rota that reproduces the real sheet's quirks.

Run with: python3 tests/fixtures/make-synthetic.py   (needs openpyxl)
"""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

DAYS = ["Δευτέρα", "Τρίτη", "Τετάρτη", "Πέμπτη", "Παρασκευή", "Σάββατο ", "Κυριακή"]
ROSTER = ["Αλεξίου Ν.", "Βλάχου Ρ.", "Ζαφειρίου Θ.", "Δημητρίου Λ.", "Στεφανάκη Μ.", "Καραλής Π.", "Ορφανού Ε.", "Βλάχου Ρ."]


def week(ws, row, dates, body):
    for i, name in enumerate(DAYS):
        ws.cell(row, i + 1, name)
    for i, value in enumerate(dates):
        cell = ws.cell(row + 1, i + 1, value)
        if isinstance(value, datetime):
            cell.number_format = "d/m/yyyy"
    r = row + 2
    for line in body:
        if line is None:
            r += 1
            continue
        names, time, *notes = line
        for i, name in enumerate(names):
            if name:
                ws.cell(r, i + 1, name)
        if time is not None:
            ws.cell(r, 8, time)
        for j, note in enumerate(notes):
            ws.cell(r, 9 + j, note)
        r += 1
    return r


wb = Workbook()
draft = wb.active
draft.title = "Sheet1"
# An old draft of the 14/9 week; Sheet2's version must win.
week(draft, 1, ["14/9", "15/9", "16/9", "17/9", "18/9", "19/9", "20/9"], [
    (["Ορφανού Ε.", "", "", "", "", "", ""], "08:00-16:00"),
])

ws = wb.create_sheet("Sheet2")
for i, name in enumerate(ROSTER):
    ws.cell(1, i + 1, name)

# Week of 14/9/2026: typed dates with every typo seen in the wild. Sunday cell is wrong on purpose.
r = week(ws, 2, ["14/9/2026", "15/92026", "16/9//2026", "17/9/2026", "18/09/2026", "19/9", "31/8/2026"], [
    (["Αλεξίου Ν.", "Αλεξίου Ν.", "Βλάχου Ρ.", "Zαφειρίου Θ.", "Αλεξίου Ν. (08:00-15:00)", "Καραλής Π.", ""], "08:00-16:00"),
    (["Δημητρίου Λ.", "Ζαφειρίου Θ.", "Αλεξίου Ν.", "Βλάχου Ρ.", "Ζαφειρίου Θ.", "", ""], "13:00-21:00"),
    (["Στεφανάκη", "Στεφανάκη Μ.", "", "Στεφανάκη Μ.", "", "", ""], None),
    None,
    (["", "", "", "", "", "Βλάχου Ρ.", "Ζαφειρίου Θ."], "16:00-00:00", "ΕΦΗΜΕΡΙΑ"),
    (["", "", "", "", "", "Δημητριου Λ", ""], "12:00-20:01"),
    (["Βλάχου Ρ.", "Δημητρίου Λ.", "Ζαφειρίου Θ.", "Αλεξίου Ν.", "Δημητρίου Λ.", "Αλεξίου Ν.", ""], "ρεπο"),
    (["Καραλής Π.", "", "Στεφανάκη Μ.", "Καραλής Π.", "Καραλής Π.", "Στεφανάκη Μ.", ""], "ρεπο"),
    (["Ορφανού Ε.", "Ορφανού Ε.", "Ορφανού Ε.", "", "", "", ""], "ΑΔΕΙΑ", 2025),
    (["", "", "", "", "Βλάχου Ρ.", "", ""], "ΓΟΝΙΚΗ ΑΔΕΙΑ"),
])

# Week of 21/9/2026: real date cells.
r = week(ws, r + 2, [datetime(2026, 9, 21 + i) for i in range(7)], [
    (["Βλάχου Ρ.", "ΚΑΡΑΛΗΣ", "", "", "", "", ""], "08:00-16:00"),
    (["", "", "", "", "", "", "Αλεξίου Ν."], "08:00-16:00", "ΕΦΗΜΕΡΙΑ"),
    (["Αλεξίου Ν.", "", "", "", "", "", ""], "ρεπο"),
])

# Week of 28/9: no year anywhere; must borrow 2026 from its neighbours.
for i, name in enumerate(ROSTER):
    ws.cell(r + 1, i + 1, name)
r = week(ws, r + 2, ["28/09", "29/09", "30/09", "01/10", "02/10", "03/10", "04/10"], [
    (["Ζαφειρίου Θ.", "", "", "", "", "", ""], "09:00-17:00"),
])

# Week of 5/10/2026 typed "5/10", parsed month-first by Sheets (5 Oct → 10 May).
swapped = [datetime(2026, 5 + i, 10) for i in range(7)]
r = week(ws, r + 2, swapped, [
    (["Αλεξίου Ν.", "", "", "", "", "", ""], "08:00-16:00"),
])

# The same week again, corrected: the later block must win.
week(ws, r + 2, swapped, [
    (["Αλεξίου Ν.", "", "", "", "", "", ""], "13:00-21:00"),
])

out = Path(__file__).with_name("synthetic.xlsx")
wb.save(out)
print(f"wrote {out}")
